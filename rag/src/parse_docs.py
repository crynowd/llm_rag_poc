import json
import os
from dataclasses import dataclass
from typing import Any, Dict, Iterable, List, Optional

from tqdm import tqdm

# PDF
from pypdf import PdfReader

# DOCX
from docx import Document as DocxDocument

import logging
from datetime import datetime

@dataclass
class DocSpec:
    doc_id: str
    title: str
    type: str  # "pdf" | "docx"
    path: str
    edition_date: str


def load_config(config_path: str) -> Dict[str, Any]:
    with open(config_path, "r", encoding="utf-8") as f:
        return json.load(f)


def ensure_dir(path: str) -> None:
    os.makedirs(path, exist_ok=True)

def make_run_dir(project_dir: str) -> str:
    run_id = datetime.now().strftime("%Y%m%d_%H%M%S")
    run_dir = os.path.join(project_dir, "runs", run_id)
    ensure_dir(run_dir)
    return run_dir


def setup_logger(log_path: str) -> logging.Logger:
    logger = logging.getLogger("rag_poc")
    logger.setLevel(logging.INFO)
    logger.handlers.clear()

    fmt = logging.Formatter(
        fmt="%(asctime)s | %(levelname)s | %(message)s",
        datefmt="%Y-%m-%d %H:%M:%S",
    )

    # File handler
    fh = logging.FileHandler(log_path, encoding="utf-8")
    fh.setLevel(logging.INFO)
    fh.setFormatter(fmt)

    # Console handler
    ch = logging.StreamHandler()
    ch.setLevel(logging.INFO)
    ch.setFormatter(fmt)

    logger.addHandler(fh)
    logger.addHandler(ch)
    return logger

def write_jsonl(path: str, rows: Iterable[Dict[str, Any]]) -> None:
    with open(path, "w", encoding="utf-8") as f:
        for r in rows:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")


def file_exists(path: str) -> bool:
    return os.path.isfile(path)


def parse_pdf(spec: DocSpec) -> List[Dict[str, Any]]:
    if not file_exists(spec.path):
        raise FileNotFoundError(f"PDF not found: {spec.path}")

    reader = PdfReader(spec.path)
    out: List[Dict[str, Any]] = []

    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        text = text.replace("\u00a0", " ").strip()  # NBSP -> space
        is_empty = 1 if not text else 0
        out.append(
            {
                "doc_id": spec.doc_id,
                "doc_title": spec.title,
                "doc_type": "pdf",
                "edition_date": spec.edition_date,
                "source_path": spec.path,
                "page": i,
                "para_idx": None,
                "text": text,
                "is_empty": is_empty,
            }
        )

    return out


def _iter_docx_paragraph_texts(doc: DocxDocument) -> Iterable[str]:
    # paragraphs
    for p in doc.paragraphs:
        t = (p.text or "").replace("\u00a0", " ").strip()
        if t:
            yield t

    # tables -> rows joined by " | " (simple v1)
    for table in doc.tables:
        for row in table.rows:
            cells = []
            for cell in row.cells:
                ct = (cell.text or "").replace("\u00a0", " ").strip()
                ct = " ".join(ct.split())
                cells.append(ct)
            line = " | ".join(cells).strip(" |")
            if line:
                yield line


def parse_docx(spec: DocSpec) -> List[Dict[str, Any]]:
    if not file_exists(spec.path):
        raise FileNotFoundError(f"DOCX not found: {spec.path}")

    doc = DocxDocument(spec.path)
    out: List[Dict[str, Any]] = []

    para_idx = 0
    for text in _iter_docx_paragraph_texts(doc):
        para_idx += 1
        out.append(
            {
                "doc_id": spec.doc_id,
                "doc_title": spec.title,
                "doc_type": "docx",
                "edition_date": spec.edition_date,
                "source_path": spec.path,
                "page": None,
                "para_idx": para_idx,
                "text": text,
                "is_empty": 0,
            }
        )

    return out


def main() -> None:
    base_dir = os.path.dirname(os.path.abspath(__file__))
    project_dir = os.path.abspath(os.path.join(base_dir, ".."))
    config_path = os.path.join(project_dir, "config.json")

    run_dir = make_run_dir(project_dir)
    log_path = os.path.join(run_dir, "parse.log")
    logger = setup_logger(log_path)

    logger.info("=== PARSE START ===")
    logger.info(f"project_dir={project_dir}")
    logger.info(f"config_path={config_path}")
    logger.info(f"run_dir={run_dir}")

    cfg = load_config(config_path)

    docs: List[DocSpec] = []
    for d in cfg.get("documents", []):
        docs.append(
            DocSpec(
                doc_id=d["doc_id"],
                title=d["title"],
                type=d["type"].lower(),
                path=d["path"],
                edition_date=d["edition_date"],
            )
        )

    logger.info(f"documents_count={len(docs)}")
    for spec in docs:
        logger.info(f"doc: doc_id={spec.doc_id} type={spec.type} edition_date={spec.edition_date} path={spec.path}")

    out_dir = os.path.join(project_dir, "data", "parsed")
    ensure_dir(out_dir)
    out_path = os.path.join(out_dir, "parsed.jsonl")

    parsed_rows: List[Dict[str, Any]] = []

    logger.info("Parsing documents...")
    per_doc_stats = []

    for spec in tqdm(docs, desc="Parsing documents"):
        if spec.type == "pdf":
            rows = parse_pdf(spec)
        elif spec.type == "docx":
            rows = parse_docx(spec)
        else:
            raise ValueError(f"Unsupported doc type: {spec.type} (doc_id={spec.doc_id})")

        parsed_rows.extend(rows)

        # per-doc stats
        total = len(rows)
        empty = sum(1 for r in rows if int(r.get("is_empty", 0)) == 1)
        nonempty = total - empty
        avg_len = int(sum(len(r.get("text", "")) for r in rows) / max(total, 1))

        per_doc_stats.append((spec.doc_id, spec.type, total, empty, nonempty, avg_len))
        logger.info(
            f"parsed doc_id={spec.doc_id} type={spec.type} segments={total} empty={empty} nonempty={nonempty} avg_chars={avg_len}"
        )

    write_jsonl(out_path, parsed_rows)

    # global stats
    total = len(parsed_rows)
    pdf_pages = sum(1 for r in parsed_rows if r["doc_type"] == "pdf")
    docx_segs = sum(1 for r in parsed_rows if r["doc_type"] == "docx")
    empty_total = sum(1 for r in parsed_rows if int(r.get("is_empty", 0)) == 1)
    avg_chars_total = int(sum(len(r.get("text", "")) for r in parsed_rows) / max(total, 1))

    logger.info(f"[OK] Written: {out_path}")
    logger.info(
        f"[STATS] total_segments={total}, pdf_pages={pdf_pages}, docx_segments={docx_segs}, empty_segments={empty_total}, avg_chars={avg_chars_total}"
    )
    logger.info("=== PARSE END ===")

    # keep existing prints (handy)
    print(f"[OK] Written: {out_path}")
    print(f"[STATS] total_segments={total}, pdf_pages={pdf_pages}, docx_segments={docx_segs}")
    print(f"[LOG] {log_path}")

if __name__ == "__main__":
    main()